import os
from typing import List
from langchain_core.documents import Document

import pypdf
from docx import Document as DocxDocument
import requests
from bs4 import BeautifulSoup

#Generate doc with text
def text_to_doc(text: str, source:str = 'inline') -> Document:
    return Document(page_content = text, metadata = {'source': source})

def load_text(path: str) -> List[Document]:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()        
    return [text_to_doc(text, source=os.path.abspath(path))] 

def load_pdf(path:str) -> List[Document]:
    rd = pypdf.PdfReader(path)
    docs = []
    for i, r in enumerate(rd.pages):
        text = r.extract_text() or ''
        print('type of the text',type(text))
        docs.append(text_to_doc(text, source = f'{os.path.abspath(path)}:page: {i+1}'))

    return docs

def load_docx(path: str) -> List[Document]:
    doc = DocxDocument(path)
    para = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    text = '\n\n'.join(para)
    return [text_to_doc(text, source = os.path.abspath(path))]

def load_url(url: str) -> List[Document]:
    response = requests.get(url, timeout = 15)
    response.raise_for_status()
    soup = BeautifulSoup(response.text)

    element = soup.find_all(['p', 'h1', 'h2', 'h3', 'li'])
    text = '\n\n'.join(e.get_text(separator=' ', strip=True) for e in element if e.get_text(strip=True))
    return [text_to_doc(text, source=url)]

def load_document(path_or_url: str) -> List[Document]:
    if path_or_url.startswith('http://') or path_or_url.startswith('https://'):
        return load_url(path_or_url) 
    if not os.path.exists(path_or_url):
        raise FileNotFoundError(f'not found: {path_or_url}')

    ex = os.path.splitext(path_or_url)[1].lower()

    if ex in ['.txt', '.md']:
        return load_text(path_or_url)
    elif ex == '.pdf':
        return load_pdf(path_or_url)
    elif ex in ['.docx', '.docs', '.doc']:
        return load_docx(path_or_url)
    else:
        try:
            return load_text(path_or_url)
        except Exception:
            raise ValueError(f'Unsupported file type: {ex}')
